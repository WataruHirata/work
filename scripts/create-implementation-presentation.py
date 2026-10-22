from pathlib import Path
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "みのり屋_実装思想_7分説明.pptx"
ASSETS = ROOT / "presentation-assets"

W = 13.333
H = 7.5

NAVY = "0A2E4E"
NAVY_2 = "123B5F"
INK = "102A43"
MUTED = "5D7083"
GREEN = "188268"
GREEN_2 = "2E9B7F"
MINT = "E3F3ED"
SKY = "EAF3F8"
CREAM = "F7F5F0"
WHITE = "FFFFFF"
LINE = "D7E1E6"
ORANGE = "E18A3B"
RED = "C95D59"
FONT = "BIZ UDPGothic"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def add_rect(slide, x, y, w, h, fill, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=20,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    line_spacing=1.05,
    font=FONT,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, runs, x, y, w, h, size=20, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.05
    for spec in runs:
        run = p.add_run()
        run.text = spec[0]
        run.font.name = FONT
        run.font.size = Pt(spec[1] if len(spec) > 1 else size)
        run.font.bold = spec[2] if len(spec) > 2 else False
        run.font.color.rgb = rgb(spec[3] if len(spec) > 3 else color)
    return box


def add_circle(slide, x, y, diameter, fill, text=None, text_color=WHITE, size=18):
    c = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(fill)
    c.line.fill.background()
    if text is not None:
        tf = c.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = rgb(text_color)
    return c


def add_pill(slide, text, x, y, w, fill=MINT, color=GREEN, size=12):
    add_rect(slide, x, y, w, 0.38, fill, None, True)
    add_text(
        slide,
        text,
        x,
        y + 0.01,
        w,
        0.34,
        size=size,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_header(slide, number, eyebrow, title, subtitle=None, dark=False):
    title_color = WHITE if dark else NAVY
    sub_color = "C7D8E5" if dark else MUTED
    add_text(slide, f"{number:02d}", 0.6, 0.42, 0.45, 0.3, 11, GREEN_2, True)
    add_text(slide, eyebrow.upper(), 1.08, 0.41, 3.8, 0.3, 10, GREEN_2, True)
    add_text(slide, title, 0.6, 0.84, 11.9, 0.62, 28, title_color, True)
    if subtitle:
        add_text(slide, subtitle, 0.62, 1.47, 11.5, 0.38, 12.5, sub_color)


def add_footer(slide, text, dark=False):
    color = "9FB6C7" if dark else "8394A3"
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(7.13), Inches(12.15), Inches(0.01)
    ).fill.solid()
    line = slide.shapes[-1]
    line.fill.fore_color.rgb = rgb("36546B" if dark else LINE)
    line.line.fill.background()
    add_text(slide, text, 0.62, 7.18, 11.5, 0.18, 8.5, color)
    add_text(slide, "みのり屋 販売管理システム", 10.6, 7.18, 2.1, 0.18, 8.5, color, align=PP_ALIGN.RIGHT)


def add_note(slide, note):
    frame = slide.notes_slide.notes_text_frame
    frame.text = note


def add_picture_cover(slide, path, x, y, w, h, alt):
    with Image.open(path) as im:
        iw, ih = im.size
    target = w / h
    source = iw / ih
    if source > target:
        shown = ih * target
        crop = (iw - shown) / 2 / iw
        pic = slide.shapes.add_picture(
            str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h)
        )
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        shown = iw / target
        crop = (ih - shown) / 2 / ih
        pic = slide.shapes.add_picture(
            str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h)
        )
        pic.crop_top = crop
        pic.crop_bottom = crop
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def add_card_title(slide, number, title, body, x, y, w, accent):
    add_rect(slide, x, y, w, 2.3, WHITE, LINE, True)
    add_circle(slide, x + 0.28, y + 0.27, 0.48, accent, str(number), WHITE, 14)
    add_text(slide, title, x + 0.28, y + 0.9, w - 0.56, 0.4, 19, NAVY, True)
    add_text(slide, body, x + 0.28, y + 1.42, w - 0.56, 0.55, 12.5, MUTED, False, line_spacing=1.18)


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(NAVY)

    add_pill(slide, "7 MINUTES / IMPLEMENTATION STORY", 0.72, 0.62, 2.95, "164568", "BFE8D9", 10)
    add_text(slide, "機能ではなく、\n判断の積み重ねを話す。", 0.72, 1.52, 7.25, 1.55, 34, WHITE, True, line_spacing=0.94)
    add_text(
        slide,
        "みのり屋 販売管理システム｜実装思想と設計の要点",
        0.76,
        3.34,
        6.6,
        0.45,
        15,
        "C7D8E5",
    )
    add_text(slide, "2026.10.22", 0.76, 6.62, 2.2, 0.28, 10, "9FB6C7", True)

    # Layered architecture motif
    for i, (x, y, w, h, fill, label) in enumerate(
        [
            (8.20, 1.04, 4.15, 4.95, "123B5F", "UI / ROUTES"),
            (8.55, 1.48, 3.45, 4.05, "18516B", "SERVER"),
            (8.90, 1.92, 2.75, 3.15, "1B6B70", "DOMAIN"),
            (9.25, 2.36, 2.05, 2.25, GREEN, "DB / IF"),
        ]
    ):
        add_rect(slide, x, y, w, h, fill, None, True)
        add_text(slide, label, x + 0.2, y + 0.16, w - 0.4, 0.24, 9.5, WHITE, True, align=PP_ALIGN.CENTER)
    add_circle(slide, 9.91, 3.08, 0.72, CREAM, "守", NAVY, 20)
    add_text(slide, "業務の正確さ", 9.32, 4.14, 1.9, 0.3, 10.5, WHITE, True, align=PP_ALIGN.CENTER)

    add_note(
        slide,
        "0:00–0:35\n今日は機能一覧ではなく、何を守るためにどう実装したかを7分で説明します。"
        "結論は、速く機能を足す前に、壊してはいけない境界を決めた、です。",
    )
    return slide


def slide_principles(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(CREAM)
    add_header(slide, 1, "starting point", "要件を足す前に、境界を決めた", "112店舗の業務を止めずに、正確さと将来変更を両立する")

    add_card_title(slide, 1, "置き換えない", "既存POS・基幹は残す。\nWebは統合・承認・可視化に集中。", 0.65, 2.25, 3.83, GREEN)
    add_card_title(slide, 2, "画面を信頼しない", "金額・権限・承認は\nServer Action と DB で再検査。", 4.75, 2.25, 3.83, NAVY_2)
    add_card_title(slide, 3, "未確定を作り込まない", "Mock → DB → 外部IF。\n決まった境界から段階接続。", 8.85, 2.25, 3.83, ORANGE)

    add_rect(slide, 1.40, 5.27, 10.53, 1.08, NAVY, None, True)
    add_text(slide, "設計の軸", 1.72, 5.56, 1.15, 0.34, 12, "9FDAC8", True)
    add_text(slide, "継続性 × 正確性 × 変更容易性", 3.05, 5.44, 7.95, 0.5, 22, WHITE, True, align=PP_ALIGN.CENTER)

    add_footer(slide, "根拠: 統合要件書v3 / ADR-0001 / docs/architecture.md")
    add_note(
        slide,
        "0:35–1:25\n最初に決めたのは3つです。既存POSを全面置換しない。重要操作はブラウザ表示だけで守らない。"
        "そして、外部IFなど未確定事項を暗黙に完成扱いしない。この3つを、継続性・正確性・変更容易性という設計軸にしました。",
    )


def slide_people(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_header(slide, 2, "experience design", "利用者ごとに、入口と責任を分ける", "共用PCの速さと、重要操作の個人識別を両立")

    # Left lane
    add_rect(slide, 0.65, 2.02, 5.85, 4.45, SKY, None, True)
    add_pill(slide, "STORE DEVICE", 0.98, 2.30, 1.62, WHITE, NAVY, 10)
    add_text(slide, "店舗端末", 0.98, 2.82, 2.2, 0.38, 21, NAVY, True)
    add_text(slide, "すぐ見る", 0.98, 3.33, 2.05, 0.48, 26, GREEN, True)
    add_text(slide, "自店舗の状況・通知\n端末スコープで閲覧", 0.98, 3.98, 2.25, 0.85, 13, MUTED, line_spacing=1.2)
    add_picture_cover(slide, ASSETS / "store.png", 3.28, 2.38, 2.84, 3.56, "店舗用共通メニューの実画面")
    add_rect(slide, 3.28, 2.38, 2.84, 3.56, "FFFFFF", LINE, True).fill.transparency = 100

    # Right lane
    add_rect(slide, 6.82, 2.02, 5.85, 4.45, MINT, None, True)
    add_pill(slide, "PERSONAL ACCOUNT", 7.15, 2.30, 1.92, WHITE, GREEN, 10)
    add_text(slide, "個人ログイン", 7.15, 2.82, 2.35, 0.38, 21, NAVY, True)
    add_text(slide, "責任を残す", 7.15, 3.33, 2.05, 0.48, 26, GREEN, True)
    add_text(slide, "保存・金額・承認\n操作者と権限を記録", 7.15, 3.98, 2.25, 0.85, 13, MUTED, line_spacing=1.2)
    add_picture_cover(slide, ASSETS / "login.png", 9.45, 2.38, 2.84, 3.56, "管理機能ログインの実画面")
    add_rect(slide, 9.45, 2.38, 2.84, 3.56, "FFFFFF", LINE, True).fill.transparency = 100

    add_rect(slide, 4.76, 6.18, 3.82, 0.50, NAVY, None, True)
    add_text(slide, "低摩擦の閲覧 / 追跡可能な更新", 4.78, 6.27, 3.78, 0.28, 12, WHITE, True, align=PP_ALIGN.CENTER)
    add_footer(slide, "ルート設計: 店舗5機能 / 管理10機能 / 機能単位の権限制御")
    add_note(
        slide,
        "1:25–2:15\n店舗の共用PCでは、閲覧まで毎回個人ログインを要求すると現場負荷が上がります。"
        "そこで端末接続は自店舗の閲覧に限定。一方、保存・金額・承認は個人ログインを必須にし、誰が操作したかを残します。",
    )


def slide_workflow(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(CREAM)
    add_header(slide, 3, "business integrity", "重要操作は、画面ではなく業務フローで守る", "例: 売上修正は「申請」と「数値更新」を分離")

    labels = [
        ("01", "店舗", "修正を申請", GREEN),
        ("02", "エリア", "一次承認", NAVY_2),
        ("03", "本部", "最終承認", NAVY),
        ("04", "DB", "売上更新・通知", ORANGE),
    ]
    xs = [0.72, 3.72, 6.72, 9.72]
    for i, ((num, actor, action, color), x) in enumerate(zip(labels, xs)):
        add_rect(slide, x, 2.35, 2.35, 1.75, WHITE, LINE, True)
        add_circle(slide, x + 0.18, 2.54, 0.44, color, num, WHITE, 10)
        add_text(slide, actor, x + 0.72, 2.56, 1.34, 0.30, 12, MUTED, True)
        add_text(slide, action, x + 0.22, 3.17, 1.92, 0.45, 18, NAVY, True, align=PP_ALIGN.CENTER)
        if i < 3:
            add_text(slide, "›", x + 2.47, 2.93, 0.38, 0.45, 28, GREEN_2, True, align=PP_ALIGN.CENTER)

    add_rect(slide, 2.03, 4.52, 9.28, 0.84, "FBEAE7", None, True)
    add_text(slide, "最終承認までは、売上を変更しない", 2.18, 4.72, 8.98, 0.38, 20, RED, True, align=PP_ALIGN.CENTER)

    guards = [("ROLE", "誰ができるか"), ("SCOPE", "どの店舗か"), ("AUDIT", "何が変わったか")]
    for i, (tag, body) in enumerate(guards):
        x = 2.28 + i * 3.0
        add_pill(slide, tag, x, 5.86, 0.90, NAVY, WHITE, 9)
        add_text(slide, body, x + 1.04, 5.90, 1.72, 0.30, 12.5, INK, True)

    add_footer(slide, "Server Action → Repository → PostgreSQL RPC / RLS / 監査ログ")
    add_note(
        slide,
        "2:15–3:15\n高リスクな売上修正は、フォームを作って終わりではありません。店舗申請、エリア一次承認、本部最終承認を状態遷移として実装し、"
        "最終承認までは売上を変えません。さらにDB側で役割・担当範囲・監査を強制し、画面の迂回でも壊れないようにしています。",
    )


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(NAVY)
    add_header(slide, 4, "changeability", "接続先を替えても、業務は替えない", "階層型モノリス + Repository / Adapter", dark=True)

    layers = [
        ("UI / ROUTES", "画面・URL", 1.0, 2.10, 3.10, "174362"),
        ("SERVER", "入力・認証・手順", 1.0, 3.12, 3.10, "18516B"),
        ("DOMAIN", "業務計算・契約", 1.0, 4.14, 3.10, "1B6B70"),
    ]
    for label, body, x, y, w, fill in layers:
        add_rect(slide, x, y, w, 0.78, fill, None, True)
        add_text(slide, label, x + 0.22, y + 0.14, 1.24, 0.25, 10, "9FDAC8", True)
        add_text(slide, body, x + 1.42, y + 0.12, 1.42, 0.30, 14, WHITE, True, align=PP_ALIGN.RIGHT)

    add_rect(slide, 4.65, 2.10, 2.00, 2.82, GREEN, None, True)
    add_text(slide, "CONTAINER", 4.90, 2.42, 1.50, 0.26, 10, "D5F1E8", True, align=PP_ALIGN.CENTER)
    add_text(slide, "接続実装を\n1か所で選ぶ", 4.90, 3.12, 1.50, 0.82, 19, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, line_spacing=1.1)

    adapters = [
        ("MOCK", "画面・業務の先行検証", "EAF3F8", NAVY),
        ("MYSQL", "ローカル接続の再現", "E3F3ED", GREEN),
        ("SUPABASE", "Auth / RPC / RLS", "F7F5F0", NAVY),
    ]
    for i, (name, desc, fill, color) in enumerate(adapters):
        y = 2.10 + i * 1.26
        add_rect(slide, 7.22, y, 2.15, 0.90, fill, None, True)
        add_text(slide, name, 7.46, y + 0.17, 1.66, 0.24, 11, color, True)
        add_text(slide, desc, 7.46, y + 0.49, 1.66, 0.22, 9.5, MUTED)

    add_rect(slide, 10.05, 2.10, 2.25, 3.42, "123B5F", "36546B", True)
    add_text(slide, "EXTERNAL IF", 10.32, 2.38, 1.71, 0.28, 10, "9FDAC8", True, align=PP_ALIGN.CENTER)
    add_text(slide, "既存POS\n基幹システム", 10.32, 3.08, 1.71, 0.75, 18, WHITE, True, align=PP_ALIGN.CENTER)
    add_pill(slide, "OUTBOX / RETRY", 10.38, 4.45, 1.58, GREEN, WHITE, 8.5)
    add_text(slide, "設定済み接続先だけ送信", 10.25, 4.98, 1.85, 0.25, 9.5, "C7D8E5", align=PP_ALIGN.CENTER)

    add_text(slide, "→", 4.21, 3.15, 0.35, 0.45, 25, "9FDAC8", True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 6.73, 3.15, 0.35, 0.45, 25, "9FDAC8", True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 9.53, 3.15, 0.35, 0.45, 25, "9FDAC8", True, align=PP_ALIGN.CENTER)

    add_rect(slide, 1.15, 6.03, 11.05, 0.64, "143E5C", None, True)
    add_text(slide, "業務契約は固定　／　接続実装は交換可能", 1.35, 6.18, 10.65, 0.28, 16, WHITE, True, align=PP_ALIGN.CENTER)
    add_footer(slide, "src/domain → src/server/container.ts → infrastructure", dark=True)
    add_note(
        slide,
        "3:15–4:20\n将来DBや外部接続が変わっても、業務ルールまで書き換えない構成にしました。domainはNext.jsやSupabaseを知らず、"
        "containerがMock、MySQL、Supabaseを選びます。外部送信はOutbox、冪等キー、再試行で失敗を業務処理から切り離しています。",
    )


def slide_ux(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_header(slide, 5, "operational ux", "現場で迷わせないUI", "見た目の統一より、判断と復帰のしやすさを優先")

    add_rect(slide, 0.70, 2.04, 8.45, 4.62, "F4F7F8", LINE, True)
    add_picture_cover(slide, ASSETS / "store.png", 0.82, 2.16, 8.21, 4.38, "店舗共通メニューの実画面")

    callouts = [
        ("01", "常に同じ場所", "共通サイドバー", 9.52, 2.13, NAVY),
        ("02", "行動で命名", "店舗業務 / 特売・価格", 9.52, 3.46, GREEN),
        ("03", "状態を明示", "SAMPLE・未接続・権限不足", 9.52, 4.79, ORANGE),
    ]
    for num, title, body, x, y, color in callouts:
        add_circle(slide, x, y, 0.44, color, num, WHITE, 10)
        add_text(slide, title, x + 0.62, y - 0.01, 2.45, 0.30, 15, NAVY, True)
        add_text(slide, body, x + 0.62, y + 0.35, 2.45, 0.28, 10.5, MUTED)

    add_pill(slide, "KEYBOARD", 9.52, 6.02, 1.05, SKY, NAVY, 8.5)
    add_pill(slide, "REDUCED MOTION", 10.68, 6.02, 1.45, MINT, GREEN, 8.5)
    add_footer(slide, "店舗共用PC / フォーカス / エラー回復 / 表示倍率 / prefers-reduced-motion")
    add_note(
        slide,
        "4:20–5:10\nUIでは、情報量を増やすより、入口と状態を揃えました。共通サイドバー、業務の言葉で付けた画面名、"
        "サンプル・未接続・権限不足の明示です。キーボード操作や動きの低減も、店舗の共用PCを前提にしています。",
    )


def slide_quality(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(CREAM)
    add_header(slide, 6, "evidence", "実装完了を、テストで言語化する", "画面の成功だけでなく、境界・業務・追跡を検証")

    metrics = [
        ("19", "test files", "構造・認可・DB・連携", NAVY),
        ("78", "tests", "Vitest 全件成功", GREEN),
        ("5", "traceability", "要件ID・版番号の整合", ORANGE),
    ]
    for i, (value, label, body, color) in enumerate(metrics):
        x = 0.72 + i * 4.18
        add_rect(slide, x, 2.08, 3.75, 2.05, WHITE, LINE, True)
        add_text(slide, value, x + 0.26, 2.40, 1.28, 0.68, 36, color, True)
        add_text(slide, label.upper(), x + 1.58, 2.52, 1.78, 0.28, 10, MUTED, True, align=PP_ALIGN.RIGHT)
        add_text(slide, body, x + 0.28, 3.38, 3.18, 0.34, 13, NAVY, True)

    add_text(slide, "次のゲート", 0.78, 4.70, 1.40, 0.32, 12, RED, True)
    add_rect(slide, 2.12, 4.55, 10.48, 1.31, "FBEAE7", None, True)
    add_text(slide, "本番IFの合意  /  実アカウントE2E  /  負荷・復旧条件", 2.45, 4.82, 9.84, 0.38, 18, RED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "未確定を「完成」と呼ばないことも、品質の一部。", 2.45, 5.31, 9.84, 0.28, 11.5, MUTED, align=PP_ALIGN.CENTER)

    add_rect(slide, 1.78, 6.23, 9.78, 0.58, NAVY, None, True)
    add_text(slide, "速く作るために、先に壊れにくくした。", 2.05, 6.34, 9.24, 0.32, 18, WHITE, True, align=PP_ALIGN.CENTER)
    add_footer(slide, "2026-10-22 検証: npm test / typecheck / Python traceability")
    add_note(
        slide,
        "5:10–6:35\n品質はテスト件数だけでなく、何を契約として守るかです。構造境界、承認や権限などの業務、要件IDとマイグレーションの追跡を自動検証しています。"
        "一方、本番IF、実アカウントE2E、負荷・復旧は次のゲートです。未確定を完成と呼ばないことも品質だと考えています。\n"
        "6:35–7:00\nまとめると、速く作るために先に壊れにくくした。これが今回の実装思想です。",
    )


def validate(prs):
    errors = []
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            right = shape.left + shape.width
            bottom = shape.top + shape.height
            if shape.left < 0 or shape.top < 0 or right > prs.slide_width + 2000 or bottom > prs.slide_height + 2000:
                errors.append(f"slide {index}: out-of-bounds shape {shape.name}")
    if errors:
        raise RuntimeError("\n".join(errors))


def main():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "みのり屋 販売管理システム｜実装思想と設計の要点"
    prs.core_properties.subject = "7分間の口頭説明用プレゼンテーション"
    prs.core_properties.author = "OpenAI Codex"
    prs.core_properties.keywords = "販売管理, Next.js, Supabase, 設計思想, 実装説明"

    slide_cover(prs)
    slide_principles(prs)
    slide_people(prs)
    slide_workflow(prs)
    slide_architecture(prs)
    slide_ux(prs)
    slide_quality(prs)
    validate(prs)
    prs.save(OUT)
    print(f"created: {OUT}")
    print(f"slides: {len(prs.slides)}")


if __name__ == "__main__":
    sys.exit(main())
